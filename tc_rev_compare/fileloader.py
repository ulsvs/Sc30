"""Structure Manager'dan alınmış CSV / XLSX BOM export'unu ağaca çevirir.

Gerekli sütunlar: seviye ve Item ID. Diğerleri (revizyon, isim, find no,
ağırlık, adet, X/Y/Z [mm] veya 16 değerli dönüşüm) varsa kullanılır.
Sütun başlıkları ayar dosyasındaki [dosya_sutunlari] ile eşleştirilir.
"""

from __future__ import annotations

import csv
from pathlib import Path

from .config import Config, parse_float
from .model import BomNode


def _read_rows(path: Path) -> list[list[str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    return list(csv.reader(text.splitlines(), dialect))


def _norm(s) -> str:
    return str(s).strip().casefold().replace("ı", "i").replace("i̇", "i")


def load_bom_file(path: str | Path, cfg: Config) -> BomNode:
    path = Path(path)
    rows = [r for r in _read_rows(path) if any(str(v).strip() for v in r)]
    if not rows:
        raise ValueError(f"{path}: dosya boş")
    header = [_norm(h) for h in rows[0]]
    col: dict[str, int] = {}
    for key, names in cfg.file_columns.items():
        for name in names:
            if _norm(name) in header:
                col[key] = header.index(_norm(name))
                break
    for required in ("seviye", "item_id"):
        if required not in col:
            raise ValueError(f"{path}: '{required}' sütunu bulunamadı. Başlıklar: {rows[0]}")

    def get(row, key):
        i = col.get(key)
        return row[i] if i is not None and i < len(row) else ""

    root: BomNode | None = None
    stack: list[tuple[int, BomNode]] = []
    for line_no, row in enumerate(rows[1:], start=2):
        level = parse_float(get(row, "seviye"))
        if level is None:
            raise ValueError(f"{path}:{line_no}: seviye okunamadı")
        level = int(level)
        node = BomNode(
            item_id=str(get(row, "item_id")).strip(),
            revision=str(get(row, "revizyon")).strip(),
            name=str(get(row, "isim")).strip(),
            find_no=str(get(row, "find_no")).strip(),
            weight=_scaled(parse_float(get(row, "agirlik")), cfg.weight_factor),
            count=max(parse_float(get(row, "adet")) or 1.0, 1.0),
        )
        if "donusum" in col and str(get(row, "donusum")).strip():
            node.transform = cfg.parse_transform(get(row, "donusum"))
        else:
            xyz = [parse_float(get(row, k)) for k in ("x", "y", "z")]
            if all(v is not None for v in xyz):
                node.transform = [1.0, 0.0, 0.0, 0.0, 0.0, 1.0, 0.0, 0.0,
                                  0.0, 0.0, 1.0, 0.0, xyz[0], xyz[1], xyz[2], 1.0]  # type: ignore[list-item]

        while stack and stack[-1][0] >= level:
            stack.pop()
        if not stack:
            if root is not None:
                raise ValueError(f"{path}:{line_no}: birden fazla kök satır var")
            root = node
        else:
            stack[-1][1].add(node)
        stack.append((level, node))
    assert root is not None
    return root


def _scaled(v: float | None, factor: float) -> float | None:
    return None if v is None else v * factor
