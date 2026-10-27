"""Karşılaştırma sonucunu konsola özetler ve Excel'e yazar."""

from __future__ import annotations

from pathlib import Path

from .compare import DEGISTIRILDI, EKLENDI, ESLESTI, SILINDI, CompareResult, LineDiff


def _kg(v: float | None, sign: bool = False) -> str:
    if v is None:
        return "-"
    return f"{v:+,.3f} kg" if sign else f"{v:,.3f} kg"


def _xyz(p) -> str:
    return "-" if p is None else f"({p[0]:.1f}, {p[1]:.1f}, {p[2]:.1f})"


def _what(d: LineDiff) -> str:
    if d.status != ESLESTI:
        return {EKLENDI: "EKLENDİ", SILINDI: "SİLİNDİ", DEGISTIRILDI: "DEĞİŞTİRİLDİ"}[d.status]
    parts = []
    if d.weight_changed:
        parts.append("ağırlık")
    if d.moved:
        parts.append(f"konum {d.distance_mm:.1f} mm")
    if d.rotated:
        parts.append(f"açı {d.rotation_deg:.2f}°")
    if d.name_changed:
        parts.append("isim")
    return ", ".join(parts) or "revizyon"


def summary_text(r: CompareResult, header: str = "", top: int = 10) -> str:
    lines = []
    if header:
        lines += [header, ""]
    o, n = r.old_total, r.new_total
    if o is not None or n is not None:
        delta = (n or 0.0) - (o or 0.0)
        pct = f" ({delta / o * 100:+.2f}%)" if o else ""
        lines.append(f"Toplam ağırlık : {_kg(o)}  ->  {_kg(n)}   Fark: {_kg(delta, True)}{pct}")
        rest = delta - r.explained_delta
        if abs(rest) > 0.001:
            lines.append(f"  (listelenen değişikliklerle açıklanamayan fark: {_kg(rest, True)})")
    else:
        lines.append("Toplam ağırlık : ağırlık özelliği okunamadı (ayar dosyasında [ozellikler].agirlik)")
    old_count = sum(1 for _ in r.old_root.walk())
    new_count = sum(1 for _ in r.new_root.walk())
    lines.append(f"BOM satırı     : {old_count}  ->  {new_count}")
    lines.append("")
    lines.append("Değişiklik özeti")
    lines.append(f"  Ağırlığı değişen          : {len(r.weight_changes)}")
    lines.append(f"  Konumu/açısı değişen      : {len(r.position_changes)}")
    lines.append(f"  İsmi değişen              : {len(r.name_changes)}")
    lines.append(f"  Eklenen                   : {len(r.of(EKLENDI))}")
    lines.append(f"  Silinen                   : {len(r.of(SILINDI))}")
    lines.append(f"  Değiştirilen (farklı parça): {len(r.of(DEGISTIRILDI))}")

    impacts = r.top_weight_impacts(top)
    if impacts:
        lines += ["", f"Ağırlığa en çok etki eden {len(impacts)} değişiklik"]
        for d in impacts:
            lines.append(f"  {d.weight_delta:+10.3f} kg  {_what(d):<22} {d.node.label:<22} {d.node.name[:40]}")

    if r.name_changes:
        lines += ["", "İsim değişiklikleri"]
        for d in r.name_changes[:top]:
            lines.append(f"  {d.node.item_id:<18} '{d.old.name}'  ->  '{d.new.name}'")
        if len(r.name_changes) > top:
            lines.append(f"  ... ve {len(r.name_changes) - top} tane daha (Excel raporuna bakın)")

    moved = sorted(r.position_changes, key=lambda d: d.distance_mm, reverse=True)
    if moved:
        lines += ["", "En çok yer değiştiren parçalar"]
        for d in moved[:top]:
            rot = f", {d.rotation_deg:.2f}°" if d.rotated else ""
            lines.append(f"  {d.distance_mm:9.1f} mm{rot:<10} {d.node.label:<22} {d.node.name[:40]}")
    return "\n".join(lines)


COLUMNS = [
    ("Durum", 14), ("Değişen", 26), ("Seviye", 7), ("Yol", 50),
    ("Item ID", 18), ("Eski Rev", 9), ("Yeni Rev", 9),
    ("Eski İsim", 32), ("Yeni İsim", 32),
    ("Eski Ağırlık (kg)", 14), ("Yeni Ağırlık (kg)", 14), ("Toplama Etkisi (kg)", 16),
    ("Eski X", 11), ("Eski Y", 11), ("Eski Z", 11), ("Yeni X", 11), ("Yeni Y", 11), ("Yeni Z", 11),
    ("Mesafe (mm)", 11), ("Açı (°)", 9), ("Find No", 9), ("Not", 40),
]

FILLS = {"red": "F8CBAD", "orange": "FCE4D6", "green": "E2EFDA", "yellow": "FFF2CC"}


def _row(d: LineDiff) -> list:
    o, n = d.old, d.new
    op = o.world_position if o else None
    np_ = n.world_position if n else None
    item = (f"{o.item_id} -> {n.item_id}" if d.status == DEGISTIRILDI else d.node.item_id)
    return [
        d.status, _what(d), d.node.level, d.node.path, item,
        o.revision if o else "", n.revision if n else "",
        o.name if o else "", n.name if n else "",
        (o.rolled_weight() if o else None), (n.rolled_weight() if n else None),
        round(d.weight_delta, 6) if d.counted else None,
        *_xyz_cells(op), *_xyz_cells(np_),
        round(d.distance_mm, 3) if d.distance_mm else None,
        round(d.rotation_deg, 3) if d.rotation_deg else None,
        d.node.find_no, d.note,
    ]


def _xyz_cells(p) -> list:
    return [round(v, 2) for v in p] if p else [None, None, None]


def _fill(d: LineDiff) -> str | None:
    if d.status == SILINDI:
        return "red"
    if d.status == EKLENDI:
        return "green"
    if d.status == DEGISTIRILDI:
        return "yellow"
    return "orange" if d.has_change else None


def write_excel(r: CompareResult, path: str | Path, header: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(summary_text(r, header, top=25).splitlines(), start=1):
        c = ws.cell(row=i, column=1, value=line)
        c.font = Font(name="Consolas")

    sheets = [
        ("Ağırlık", sorted([d for d in r.diffs if d.weight_changed or (d.counted and abs(d.weight_delta) > 1e-9)],
                           key=lambda d: abs(d.weight_delta), reverse=True)),
        ("Konum", sorted(r.position_changes, key=lambda d: d.distance_mm, reverse=True)),
        ("İsim", r.name_changes),
        ("Eklenen-Silinen", r.of(SILINDI) + r.of(EKLENDI) + r.of(DEGISTIRILDI)),
        ("Tüm Değişiklikler", r.changed),
    ]
    bold = Font(bold=True)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for ci, (name, width) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=1, column=ci, value=name)
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[c.column_letter].width = width
        for ri, d in enumerate(rows, start=2):
            fill = _fill(d)
            for ci, v in enumerate(_row(d), start=1):
                c = ws.cell(row=ri, column=ci, value=v)
                if fill:
                    c.fill = PatternFill("solid", fgColor=FILLS[fill])
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
    path = Path(path)
    wb.save(path)
    return path
