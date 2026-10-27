from pathlib import Path

import pytest

from tc_rev_compare.compare import DEGISTIRILDI, EKLENDI, SILINDI, Tolerances, compare
from tc_rev_compare.config import Config, parse_float
from tc_rev_compare.fileloader import load_bom_file
from tc_rev_compare.model import BomNode, rotation_angle_deg
from tc_rev_compare.report import summary_text, write_excel

ORNEK = Path(__file__).resolve().parent.parent / "ornekler"


def at(x, y, z):
    return [1, 0, 0, 0, 0, 1, 0, 0, 0, 0, 1, 0, x, y, z, 1]


def rot_z90(x=0, y=0, z=0):
    return [0, 1, 0, 0, -1, 0, 0, 0, 0, 0, 1, 0, x, y, z, 1]


def test_sample_files():
    cfg = Config()
    r = compare(load_bom_file(ORNEK / "eski.csv", cfg), load_bom_file(ORNEK / "yeni.csv", cfg))
    assert r.old_total == pytest.approx(929.95)
    assert r.new_total == pytest.approx(954.55)
    assert r.explained_delta == pytest.approx(r.new_total - r.old_total)
    assert {d.node.item_id for d in r.weight_changes} == {"100100", "100210"}
    moved = {d.node.item_id: d.distance_mm for d in r.position_changes}
    assert moved == {"100300": pytest.approx(50), "100210": pytest.approx(20)}
    assert [(d.old.name, d.new.name) for d in r.name_changes] == [("KABİN BAĞLANTI SACI", "KABİN BAĞLANTI PLAKASI")]
    assert [(d.old.item_id, d.new.item_id) for d in r.of(DEGISTIRILDI)] == [("100500", "100600")]
    assert [d.new.item_id for d in r.of(EKLENDI)] == ["100700"]
    assert r.of(SILINDI) == []


def test_duplicate_instances_matched_by_position():
    old, new = BomNode("A"), BomNode("A")
    for y in (0, 100, 200):
        old.add(BomNode("B", find_no="10", weight=1, transform=at(0, y, 0)))
    for y in (200, 0, 100.05):  # sıra değişmiş, biri 0.05 mm kaymış
        new.add(BomNode("B", find_no="10", weight=1, transform=at(0, y, 0)))
    r = compare(old, new)
    assert r.changed == []


def test_removed_subassembly_counts_full_weight():
    old, new = BomNode("A"), BomNode("A")
    sub = old.add(BomNode("S", find_no="10", transform=at(0, 0, 0)))
    sub.add(BomNode("P1", weight=2))
    sub.add(BomNode("P2", weight=3))
    new.add(BomNode("X", find_no="20", weight=1, transform=at(500, 0, 0)))
    r = compare(old, new)
    assert [d.weight_delta for d in r.of(SILINDI)] == [-5]
    assert [d.weight_delta for d in r.of(EKLENDI)] == [1]
    assert r.explained_delta == pytest.approx(-4)


def test_rotation_detected():
    assert rotation_angle_deg(at(0, 0, 0), rot_z90()) == pytest.approx(90)
    old, new = BomNode("A"), BomNode("A")
    old.add(BomNode("B", transform=at(10, 0, 0)))
    new.add(BomNode("B", transform=rot_z90(10, 0, 0)))
    r = compare(old, new)
    (d,) = r.position_changes
    assert d.rotated and not d.moved


def test_weight_tolerance():
    old, new = BomNode("A"), BomNode("A")
    old.add(BomNode("B", weight=1.0))
    new.add(BomNode("B", weight=1.0004))
    assert compare(old, new).weight_changes == []
    assert len(compare(old, new, Tolerances(weight_kg=0.0001)).weight_changes) == 1


def test_world_position_through_parent():
    root = BomNode("A")
    sub = root.add(BomNode("S", transform=at(1000, 0, 0)))
    part = sub.add(BomNode("P", transform=at(0, 50, 0)))
    assert part.world_position == (1000, 50, 0)


@pytest.mark.parametrize("text,value", [("12,5", 12.5), ("1.234,5", 1234.5), ("1,234.5", 1234.5),
                                        ("3.2 kg", 3.2), ("", None), ("abc", None), ("1.5e-05", 1.5e-05)])
def test_parse_float(text, value):
    assert parse_float(text) == value


def test_report(tmp_path):
    cfg = Config()
    r = compare(load_bom_file(ORNEK / "eski.csv", cfg), load_bom_file(ORNEK / "yeni.csv", cfg))
    text = summary_text(r)
    assert "+24.600 kg" in text
    out = write_excel(r, tmp_path / "r.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Özet", "Ağırlık", "Konum", "İsim", "Eklenen-Silinen", "Tüm Değişiklikler"]
    assert wb["Konum"].max_row == 3
